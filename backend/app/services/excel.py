"""Excel files: the question template, question upload, question export and results export."""

import math
import re
from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

QUESTION_COLUMNS = [
    ("Section", 22), ("Section Hint", 30), ("Passage Title", 18), ("Passage Text", 40), ("Question", 50),
    ("Option A", 22), ("Option B", 22), ("Option C", 22), ("Option D", 22), ("Option E", 22),
    ("Correct Answer", 15), ("Marks", 8),
]
LETTERS = ["A", "B", "C", "D", "E"]
MAX_ROWS = 3000
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
IST = timezone(timedelta(hours=5, minutes=30))  # India has no daylight saving

SAMPLE_ROWS = [
    ["Synonyms", "Choose the word closest in meaning", "", "", "Pragmatic", "Idealistic", "Practical", "Emotional", "Theoretical", "", "B", 1],
    ["Reading Comprehension", "", "Hybrid Work", "In recent years, many organizations have adopted hybrid work models...",
     "The primary purpose of the passage is to:", "Advocate permanent remote work", "Compare productivity across industries",
     "Examine trade-offs in hybrid work models", "Argue against workplace flexibility", "Prove innovation declines in remote settings", "C", 1],
    ["Reading Comprehension", "", "Hybrid Work", "", "According to the passage, hybrid work may negatively impact:",
     "Employee morale", "Compensation models", "Individual productivity", "Informal collaboration", "Office real estate costs", "D", 1],
]

HOW_TO = [
    "How to fill the Questions sheet",
    "",
    "One row per question. Do not rename or reorder the header row.",
    "Section: required. Questions appear grouped by section, in the order sections first appear.",
    'Section Hint: optional, shown under the section name (e.g. "Choose the word closest in meaning"). Fill it once per section.',
    "Passage Title + Passage Text: for reading comprehension. Put the full passage text on the FIRST question of that passage;",
    "   later questions of the same passage only need the same Passage Title.",
    "Option A to Option E: at least A and B. Fill options in order without gaps. Leave unused options empty.",
    "Correct Answer: the letter of the correct option (A, B, C, D or E).",
    "Marks: optional, defaults to 1.",
]

HEADER_FONT = Font(bold=True, color="FFFFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="FF006F53")  # Rite Tropical Rainforest


def _style_header(ws) -> None:
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"


def _to_bytes(wb: Workbook) -> bytes:
    # openpyxl turns any text starting with "=" into a live formula. Values typed by candidates (e.g. an email
    # like =HYPERLINK(...)) must open in Excel as plain text, never run.
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    cell.data_type = "s"
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def questions_workbook(rows: list[list[Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Questions"
    ws.append([name for name, _ in QUESTION_COLUMNS])
    for i, (_, width) in enumerate(QUESTION_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in rows:
        ws.append(row)
    _style_header(ws)
    for col in ("D", "E"):
        for cell in ws[col][1:]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    help_ws = wb.create_sheet("How to fill")
    help_ws.column_dimensions["A"].width = 120
    for line in HOW_TO:
        help_ws.append([line])
    help_ws["A1"].font = Font(bold=True, size=14)
    return _to_bytes(wb)


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime | date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_questions_workbook(data: bytes) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Returns (rows, errors). Rows are only meaningful when errors is empty."""
    try:
        wb = load_workbook(BytesIO(data), data_only=True, read_only=True)
    except Exception:  # noqa: BLE001 - any unreadable file gets the same friendly message
        return [], [{"row": 0, "message": "This file could not be read. Please upload an .xlsx file made from the template."}]
    ws = wb["Questions"] if "Questions" in wb.sheetnames else (wb.worksheets[0] if wb.worksheets else None)
    if ws is None:
        return [], [{"row": 0, "message": "The file has no worksheets."}]

    all_rows = list(ws.iter_rows(values_only=True))
    if not all_rows:
        return [], [{"row": 0, "message": "The sheet has no questions."}]
    col = {_cell_text(v).lower(): i for i, v in enumerate(all_rows[0]) if _cell_text(v)}
    missing = [h for h in ("Section", "Question", "Option A", "Option B", "Correct Answer") if h.lower() not in col]
    if missing:
        return [], [{"row": 1, "message": f"Missing column(s): {', '.join(missing)}. Please use the template."}]

    rows: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    for n, values in enumerate(all_rows[1:], start=2):
        def get(header: str, values: tuple = values) -> str:
            i = col.get(header.lower())
            return _cell_text(values[i]) if i is not None and i < len(values) else ""

        section, text = get("Section"), get("Question")
        letters = [get(f"Option {letter}") for letter in LETTERS]
        if not section and not text and not any(letters):
            continue  # blank row
        if len(rows) + len(errors) >= MAX_ROWS:
            break

        problems = []
        if not section:
            problems.append("Section is empty")
        if not text:
            problems.append("Question is empty")
        last_filled = max((i for i, o in enumerate(letters) if o), default=-1)
        options = letters[: last_filled + 1]
        if len(options) < 2:
            problems.append("At least Option A and Option B are needed")
        elif any(not o for o in options):
            problems.append("Options must be filled in order (no gaps)")
        letter = re.sub(r"^option\s*", "", get("Correct Answer"), flags=re.IGNORECASE).strip().upper()
        correct_index = LETTERS.index(letter) if letter in LETTERS else -1
        if correct_index < 0:
            problems.append("Correct Answer must be A, B, C, D or E")
        elif correct_index >= len(options):
            problems.append(f"Correct Answer {letter} points to an empty option")
        raw_marks = get("Marks")
        try:
            marks = 1.0 if raw_marks == "" else float(raw_marks)
        except ValueError:
            marks = math.nan
        if not (0.01 <= marks <= 100):
            problems.append("Marks must be a number from 0.01 to 100")
        passage_title, passage_text = get("Passage Title"), get("Passage Text")
        if passage_text and not passage_title:
            problems.append("Passage Text needs a Passage Title")
        if len(text) > 4000 or len(passage_text) > 20000 or any(len(o) > 1000 for o in options):
            problems.append("Text is too long")

        if problems:
            errors.append({"row": n, "message": "; ".join(problems)})
        else:
            rows.append({"row": n, "section": section, "sectionHint": get("Section Hint"), "passageTitle": passage_title,
                         "passageText": passage_text, "text": text, "options": options, "correctIndex": correct_index, "marks": marks})
    wb.close()
    if not rows and not errors:
        errors.append({"row": 0, "message": "The sheet has no questions."})
    return rows, errors


def exam_to_rows(questions: list[dict[str, Any]]) -> list[list[Any]]:
    """An exam's questions in the template layout, so they can be edited and uploaded again."""
    seen_section: set[int] = set()
    seen_passage: set[int] = set()
    out = []
    for q in questions:
        first_of_section = q["section_id"] not in seen_section
        seen_section.add(q["section_id"])
        first_of_passage = bool(q["passage_id"]) and q["passage_id"] not in seen_passage
        if q["passage_id"]:
            seen_passage.add(q["passage_id"])
        options = list(q["options"])
        out.append([
            q["section"], q["section_hint"] if first_of_section else "", q["passage_title"] or "",
            q["passage_body"] if first_of_passage else "", q["text"],
            *[options[i] if i < len(options) else "" for i in range(len(LETTERS))],
            LETTERS[q["correct_index"]], q["marks"],
        ])
    return out


def _fmt_time(d: datetime | None) -> str:
    return d.astimezone(IST).strftime("%d/%m/%Y, %I:%M:%S %p").lower() if d else ""


def results_workbook(drive: dict[str, Any], section_names: list[str], attempts: list[dict[str, Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    headers = ["Rank", "First Name", "Middle Name", "Last Name", "Student ID", "Email", "Phone Number", "Status",
               "Score", "Max Score", "Percentage", *[f"{s} (score)" for s in section_names], "Answered", "Violations",
               "Registered At", "Started At", "Submitted At", "Submitted By", "Result"]
    ws.append(headers)
    for a in attempts:
        secs = {s["name"]: s["score"] for s in (a["section_scores"] or [])}
        pct = round(a["score"] / a["max_score"] * 100, 1) if a["max_score"] else ""
        result = {"shortlisted": "Shortlisted", "not_selected": "Not selected"}.get(a["result_status"] or "", "")
        ws.append([
            a.get("rank", ""), a["first_name"], a["middle_name"], a["last_name"], a["student_code"], a["email"], a["phone"], a["status"],
            a["score"] if a["score"] is not None else "", a["max_score"] if a["max_score"] is not None else "", pct,
            *[secs.get(s, "") for s in section_names], a["answered"], a["violations"],
            _fmt_time(a["registered_at"]), _fmt_time(a["started_at"]), _fmt_time(a["submitted_at"]), a["submit_reason"] or "", result,
        ])
    for i, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(10, min(36, len(header) + 4))
    _style_header(ws)
    info = wb.create_sheet("Drive")
    for row in (["Drive", drive["name"]], ["College", drive["college"]], ["Exam", drive["exam_title"]], ["Exported", _fmt_time(datetime.now(IST))]):
        info.append(row)
    for cell in info["A"]:
        cell.font = Font(bold=True)
    info.column_dimensions["A"].width = 12
    info.column_dimensions["B"].width = 50
    return _to_bytes(wb)
