"""Admins: login, users, exams and questions, Excel, drives, monitor, results, publishing."""

from io import BytesIO

from openpyxl import load_workbook

from app import db
from app.services.excel import SAMPLE_ROWS, questions_workbook
from tests.conftest import ADMIN, auth, make_drive, make_exam, register

XLSX = {"Content-Type": "application/octet-stream"}


async def start(api, token: str) -> None:
    assert (await api.post("/api/attempt/start", headers=auth(token))).status_code == 200


# ---- login and admin users ----

async def test_admin_apis_need_login(api):
    assert (await api.get("/api/admin/drives")).status_code == 401
    assert (await api.get("/api/admin/drives", headers=auth("garbage"))).status_code == 401


async def test_wrong_password_rejected(api):
    assert (await api.post("/api/admin/login", json={"email": ADMIN["email"], "password": "nope"})).status_code == 401


async def test_add_admin_and_cannot_delete_self(api, admin):
    r = await admin.post("/api/admin/admins", json={"name": "HR", "email": "hr@test.local", "password": "short"})
    assert r.status_code == 400 and "password" in r.json()["fields"]
    assert (await admin.post("/api/admin/admins", json={"name": "HR", "email": "hr@test.local", "password": "LongEnough1"})).status_code == 201
    assert (await api.post("/api/admin/login", json={"email": "hr@test.local", "password": "LongEnough1"})).status_code == 200
    assert (await admin.delete("/api/admin/admins/1")).status_code == 400


async def test_change_password(api, admin):
    r = await admin.put("/api/admin/me/password", json={"currentPassword": "wrong", "newPassword": "Another@123"})
    assert r.status_code == 400
    assert (await admin.put("/api/admin/me/password", json={"currentPassword": ADMIN["password"], "newPassword": "Another@123"})).status_code == 204
    assert (await api.post("/api/admin/login", json={"email": ADMIN["email"], "password": "Another@123"})).status_code == 200


# ---- exams and questions ----

async def test_create_exam_section_question(admin):
    e = (await admin.post("/api/admin/exams", json={"title": "Aptitude", "durationMin": 45})).json()
    s = (await admin.post(f"/api/admin/exams/{e['id']}/sections", json={"name": "Quant"})).json()
    bad = await admin.post(f"/api/admin/exams/{e['id']}/questions", json={"sectionId": s["id"], "text": "2+2?", "options": ["3", "4"], "correctIndex": 5})
    assert bad.status_code == 400
    ok = await admin.post(f"/api/admin/exams/{e['id']}/questions", json={"sectionId": s["id"], "text": "2+2?", "options": ["3", "4", ""], "correctIndex": 1})
    assert ok.status_code == 201
    full = (await admin.get(f"/api/admin/exams/{e['id']}")).json()
    assert len(full["questions"]) == 1
    assert full["questions"][0]["options"] == ["3", "4"]
    assert full["locked"] is False


async def test_import_questions_from_template(admin):
    e = (await admin.post("/api/admin/exams", json={"title": "Imported", "durationMin": 60})).json()
    file = questions_workbook(SAMPLE_ROWS)
    r = await admin.post(f"/api/admin/exams/{e['id']}/import?mode=append", headers=XLSX, content=file)
    assert r.status_code == 200, r.text
    assert (r.json()["questions"], r.json()["sections"], r.json()["passages"]) == (3, 2, 1)
    full = (await admin.get(f"/api/admin/exams/{e['id']}")).json()
    assert [(s["name"], s["hint"]) for s in full["sections"]] == [("Synonyms", "Choose the word closest in meaning"), ("Reading Comprehension", "")]
    assert len([q for q in full["questions"] if q["passage_id"]]) == 2
    assert full["questions"][0]["correct_index"] == 1

    assert (await admin.post(f"/api/admin/exams/{e['id']}/import?mode=replace", headers=XLSX, content=file)).status_code == 200
    assert len((await admin.get(f"/api/admin/exams/{e['id']}")).json()["questions"]) == 3

    exported = await admin.get(f"/api/admin/exams/{e['id']}/questions.xlsx")
    assert exported.status_code == 200
    assert (await admin.post(f"/api/admin/exams/{e['id']}/import?mode=replace", headers=XLSX, content=exported.content)).status_code == 200


async def test_bad_sheet_reports_rows_and_imports_nothing(admin):
    e = (await admin.post("/api/admin/exams", json={"title": "Bad", "durationMin": 60})).json()
    file = questions_workbook([
        ["Synonyms", "", "", "", "Good one", "x", "y", "", "", "", "A", 1],
        ["Synonyms", "", "", "", "No answer", "x", "y", "", "", "", "F", 1],
        ["", "", "", "", "No section", "x", "", "", "", "", "A", 1],
        ["RC", "", "Missing passage", "", "Q", "x", "y", "", "", "", "B", 1],
    ])
    r = await admin.post(f"/api/admin/exams/{e['id']}/import", headers=XLSX, content=file)
    assert r.status_code == 400
    assert [x["row"] for x in r.json()["rowErrors"]] == [3, 4]
    missing = questions_workbook([SAMPLE_ROWS[0], [v if i != 2 else "Missing passage" for i, v in enumerate(SAMPLE_ROWS[0])]])
    r2 = await admin.post(f"/api/admin/exams/{e['id']}/import", headers=XLSX, content=missing)
    assert r2.status_code == 400 and "Passage Text" in r2.json()["rowErrors"][0]["message"]
    assert await db.value("SELECT count(*) FROM questions") == 0
    assert (await admin.post(f"/api/admin/exams/{e['id']}/import", headers=XLSX, content=b"not excel")).status_code == 400


async def test_template_is_valid_workbook(admin):
    r = await admin.get("/api/admin/question-template.xlsx")
    assert r.status_code == 200
    wb = load_workbook(BytesIO(r.content))
    assert wb["Questions"]["A1"].value == "Section"


async def test_locked_exam_allows_key_fixes_only(api, admin):
    ex = await make_exam()
    await make_drive(ex["exam_id"])
    token = (await register(api))["token"]
    await start(api, token)
    full = (await admin.get(f"/api/admin/exams/{ex['exam_id']}")).json()
    assert full["locked"] is True
    new_question = {"sectionId": full["sections"][0]["id"], "text": "x", "options": ["a", "b"], "correctIndex": 0}
    assert (await admin.post(f"/api/admin/exams/{ex['exam_id']}/questions", json=new_question)).status_code == 409
    assert (await admin.delete(f"/api/admin/questions/{ex['question_ids'][0]}")).status_code == 409
    q0 = full["questions"][0]
    fix = {"sectionId": q0["section_id"], "text": q0["text"], "options": q0["options"], "correctIndex": 2}
    assert (await admin.put(f"/api/admin/questions/{q0['id']}", json=fix)).status_code == 200
    fewer = {**fix, "options": q0["options"][:2], "correctIndex": 0}
    assert (await admin.put(f"/api/admin/questions/{q0['id']}", json=fewer)).status_code == 409


async def test_duplicate_exam(admin):
    ex = await make_exam()
    copy = (await admin.post(f"/api/admin/exams/{ex['exam_id']}/duplicate")).json()
    full = (await admin.get(f"/api/admin/exams/{copy['id']}")).json()
    assert len(full["questions"]) == 5 and len(full["passages"]) == 1
    assert len([q for q in full["questions"] if q["passage_id"] == full["passages"][0]["id"]]) == 2
    section_ids = {s["id"] for s in full["sections"]}
    assert all(q["section_id"] in section_ids for q in full["questions"])


async def test_sections_move_and_delete(admin):
    e = (await admin.post("/api/admin/exams", json={"title": "Order", "durationMin": 30})).json()
    a = (await admin.post(f"/api/admin/exams/{e['id']}/sections", json={"name": "A"})).json()
    b = (await admin.post(f"/api/admin/exams/{e['id']}/sections", json={"name": "B"})).json()
    assert (await admin.post(f"/api/admin/sections/{b['id']}/move", json={"direction": "up"})).status_code == 204
    names = [s["name"] for s in (await admin.get(f"/api/admin/exams/{e['id']}")).json()["sections"]]
    assert names == ["B", "A"]
    assert (await admin.delete(f"/api/admin/sections/{a['id']}")).status_code == 204
    assert (await admin.get("/api/admin/exams/abc")).status_code == 404


# ---- drives and results ----

async def test_create_drive_with_generated_code_and_time_checks(admin):
    ex = await make_exam()
    times = {"opensAt": "2026-10-12T10:00:00Z", "closesAt": "2026-10-12T09:00:00Z"}
    bad = await admin.post("/api/admin/drives", json={"name": "D", "college": "X", "examId": ex["exam_id"], **times})
    assert bad.status_code == 400 and "closesAt" in bad.json()["fields"]
    d = (await admin.post("/api/admin/drives", json={"name": "Drive", "college": "IBS Hyderabad", "examId": ex["exam_id"]})).json()
    assert d["code"].startswith("ibs-hyderabad-") and len(d["code"]) == len("ibs-hyderabad-") + 6
    assert d["status"] == "open"
    assert (await admin.post("/api/admin/drives", json={"name": "D2", "college": "Y", "examId": ex["exam_id"], "code": d["code"]})).status_code == 409
    closed = (await admin.post(f"/api/admin/drives/{d['id']}/open", json={"isOpen": False})).json()
    assert closed["status"] == "closed"


async def test_rank_export_regrade_extend_and_force_submit(api, admin):
    ex = await make_exam()
    d = await make_drive(ex["exam_id"])
    students = [await register(api) for _ in range(3)]
    for s in students:
        await start(api, s["token"])

    async def save(s, i, sel):
        r = await api.post("/api/attempt/answers", headers=auth(s["token"]), json={"answers": [{"questionId": ex["question_ids"][i], "selected": sel}]})
        assert r.status_code == 200

    await save(students[0], 0, 1)  # 1 correct
    await save(students[1], 0, 1)
    await save(students[1], 1, 0)  # 2 correct
    await save(students[2], 1, 3)  # 0 correct
    for s in students[:2]:
        await api.post("/api/attempt/submit", headers=auth(s["token"]))

    listing = (await admin.get(f"/api/admin/drives/{d['id']}/attempts")).json()
    assert listing["sections"] == ["Synonyms", "Reading"]
    assert [(a.get("rank"), a["score"], a["status"]) for a in listing["attempts"]] == [(1, 2, "submitted"), (2, 1, "submitted"), (None, None, "in_progress")]

    third = listing["attempts"][2]
    assert (await admin.post(f"/api/admin/attempts/{third['id']}/extend", json={"minutes": 10})).status_code == 204
    assert await db.value("SELECT extract(epoch FROM ends_at - started_at)::int FROM attempts WHERE id = :id", id=third["id"]) == 70 * 60
    assert (await admin.post(f"/api/admin/attempts/{third['id']}/submit")).status_code == 204

    q2 = await db.one("SELECT * FROM questions WHERE id = :id", id=ex["question_ids"][1])
    fix = {"sectionId": q2["section_id"], "text": q2["text"], "options": q2["options"], "correctIndex": 3}
    assert (await admin.put(f"/api/admin/questions/{q2['id']}", json=fix)).status_code == 200
    assert (await admin.post(f"/api/admin/drives/{d['id']}/regrade")).json()["regraded"] == 3
    listing = (await admin.get(f"/api/admin/drives/{d['id']}/attempts")).json()
    assert [(a["rank"], a["score"]) for a in listing["attempts"]] == [(1, 1), (1, 1), (1, 1)]

    detail = (await admin.get(f"/api/admin/attempts/{third['id']}")).json()
    assert detail["attempt"]["submit_reason"] == "coordinator"
    assert "token_hash" not in detail["attempt"]
    assert any(e["type"] == "time_extended" for e in detail["events"])

    xlsx = await admin.get(f"/api/admin/drives/{d['id']}/results.xlsx")
    ws = load_workbook(BytesIO(xlsx.content))["Results"]
    assert ws.max_row == 4
    assert ws.cell(1, 5).value == "Student ID"


async def test_coordinator_approves_relogin(api, admin):
    ex = await make_exam()
    d = await make_drive(ex["exam_id"])
    s = await register(api)

    async def first_row():
        return (await admin.get(f"/api/admin/drives/{d['id']}/attempts")).json()["attempts"][0]

    assert (await admin.post(f"/api/admin/attempts/{(await first_row())['id']}/allow-relogin")).status_code == 409  # not started yet
    await start(api, s["token"])
    assert (await api.post("/api/drives/test-drive/register", json=s["student"])).status_code == 409
    row = await first_row()
    assert (row["relogin_pending"], row["relogin_approved"]) == (True, False)
    assert (await admin.post(f"/api/admin/attempts/{row['id']}/allow-relogin")).status_code == 204
    row = await first_row()
    assert (row["relogin_pending"], row["relogin_approved"]) == (False, True)
    assert (await api.post("/api/drives/test-drive/register", json=s["student"])).status_code == 200
    row = await first_row()
    assert (row["relogin_pending"], row["relogin_approved"]) == (False, False)
    detail = (await admin.get(f"/api/admin/attempts/{row['id']}")).json()
    assert "admin@test.local" in next(e for e in detail["events"] if e["type"] == "relogin_allowed")["detail"]


async def test_publish_results_flow(api, admin):
    ex = await make_exam()
    d = await make_drive(ex["exam_id"])

    async def take(picks):
        s = await register(api)
        await start(api, s["token"])
        answers = [{"questionId": ex["question_ids"][i], "selected": sel} for i, sel in enumerate(picks) if sel is not None]
        if answers:
            assert (await api.post("/api/attempt/answers", headers=auth(s["token"]), json={"answers": answers})).status_code == 200
        return s

    top = await take([1, 0, 1, 2, 3])  # 5/5
    mid = await take([1, 0, 0, 0, 0])  # 2/5 = 40%
    unfinished = await take([1])
    for s in (top, mid):
        await api.post("/api/attempt/submit", headers=auth(s["token"]))

    def lookup(s, phone=None):
        return api.post("/api/drives/test-drive/result", json={"studentId": s["student"]["studentId"], "phone": phone or s["student"]["phone"]})

    r = await lookup(top)
    assert r.status_code == 404 and r.json()["published"] is False

    assert (await admin.post(f"/api/admin/drives/{d['id']}/publish", json={"nextSteps": "x"})).status_code == 400
    counts = (await admin.post(f"/api/admin/drives/{d['id']}/publish", json={"cutoffPct": 60, "nextSteps": "Interview on 20 Oct, 10 am, Hyderabad office."})).json()
    assert (counts["shortlisted"], counts["not_selected"], counts["not_submitted"]) == (1, 1, 1)

    r1 = (await lookup(top)).json()
    assert (r1["status"], r1["message"]) == ("shortlisted", "Interview on 20 Oct, 10 am, Hyderabad office.")
    assert "score" not in str(r1), "no marks are sent to students"
    r2 = (await lookup(mid)).json()
    assert r2["status"] == "not_selected" and "not been shortlisted" in r2["message"]
    assert (await lookup(unfinished)).json()["status"] == "not_submitted"
    assert (await lookup(top, phone="9000000000")).status_code == 404
    assert (await api.post("/api/drives/test-drive/result", json={"studentId": "", "phone": ""})).status_code == 400

    listing = (await admin.get(f"/api/admin/drives/{d['id']}/attempts")).json()["attempts"]
    mid_row = next(a for a in listing if a["student_code"] == mid["student"]["studentId"])
    assert (await admin.put(f"/api/admin/attempts/{mid_row['id']}/result", json={"status": "shortlisted"})).status_code == 204
    assert (await admin.post(f"/api/admin/drives/{d['id']}/publish", json={"cutoffPct": 90, "nextSteps": "Interview"})).status_code == 200
    assert (await lookup(mid)).json()["status"] == "shortlisted"
    unfinished_row = next(a for a in listing if a["student_code"] == unfinished["student"]["studentId"])
    assert (await admin.put(f"/api/admin/attempts/{unfinished_row['id']}/result", json={"status": "shortlisted"})).status_code == 409

    assert (await api.get("/api/drives/test-drive")).json()["resultsPublished"] is True
    assert (await admin.post(f"/api/admin/drives/{d['id']}/unpublish")).status_code == 204
    assert (await lookup(top)).status_code == 404


async def test_delete_candidate_allows_new_registration(api, admin):
    ex = await make_exam()
    d = await make_drive(ex["exam_id"])
    s = await register(api)
    row = (await admin.get(f"/api/admin/drives/{d['id']}/attempts")).json()["attempts"][0]
    assert (await admin.delete(f"/api/admin/attempts/{row['id']}")).status_code == 204
    assert (await api.post("/api/drives/test-drive/register", json=s["student"])).status_code == 201


async def test_drive_exam_cannot_change_after_registration(api, admin):
    ex, ex2 = await make_exam(), await make_exam()
    d = await make_drive(ex["exam_id"])
    await register(api)
    r = await admin.put(f"/api/admin/drives/{d['id']}", json={"name": "Test drive", "college": "Test College", "examId": ex2["exam_id"], "code": d["code"]})
    assert r.status_code == 409
    r = await admin.put(f"/api/admin/drives/{d['id']}", json={"name": "Renamed", "college": "Test College", "examId": ex["exam_id"], "code": d["code"]})
    assert r.status_code == 200


async def test_admin_sees_camera_photos(api, admin):
    ex = await make_exam()
    d = await make_drive(ex["exam_id"])
    s = await register(api)
    await start(api, s["token"])
    jpeg = b"\xff\xd8\xff\xe0" + bytes([7]) * 300
    r = await api.post("/api/attempt/snapshot", headers={**auth(s["token"]), "Content-Type": "image/jpeg"}, content=jpeg)
    assert r.status_code == 201
    row = (await admin.get(f"/api/admin/drives/{d['id']}/attempts")).json()["attempts"][0]
    assert row["snapshot_count"] == 1
    detail = (await admin.get(f"/api/admin/attempts/{row['id']}")).json()
    img = await admin.get(f"/api/admin/snapshots/{detail['snapshots'][0]['id']}")
    assert img.status_code == 200 and img.headers["content-type"] == "image/jpeg" and img.content == jpeg


async def test_errors_are_json_and_headers_are_secure(api):
    r = await api.get("/api/nothing-here")
    assert r.status_code == 404 and r.json() == {"error": "Not found."}
    r = await api.post("/api/drives/x/register", content=b"{not json", headers={"Content-Type": "application/json"})
    assert r.status_code == 400 and r.json() == {"error": "Invalid request."}
    health = await api.get("/api/health")
    assert health.json() == {"ok": True}
    assert health.headers["x-content-type-options"] == "nosniff"
    assert "camera=(self)" in health.headers["permissions-policy"]
